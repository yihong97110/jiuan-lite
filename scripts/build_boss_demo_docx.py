from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "64748B"
LIGHT = "F2F4F7"
CODE_FILL = "F6F8FA"
WHITE = "FFFFFF"
CONTENT_DXA = 9360


def set_run_font(run, ascii_name="Calibri", east_asia="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = ascii_name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def widths_for(count):
    patterns = {
        1: [9360],
        2: [2600, 6760],
        3: [1700, 3000, 4660],
        4: [1350, 2350, 2350, 3310],
        5: [1300, 1900, 1900, 1900, 2360],
        6: [1150, 1650, 1650, 1650, 1650, 1610],
        7: [900, 1800, 1800, 1050, 1050, 1050, 1710],
    }
    if count in patterns:
        return patterns[count]
    base = CONTENT_DXA // count
    widths = [base] * count
    widths[-1] += CONTENT_DXA - sum(widths)
    return widths


def add_inline(paragraph, text, base_size=10.5, color=INK):
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=base_size, color=color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, ascii_name="Consolas", east_asia="Microsoft YaHei", size=max(8, base_size - 0.5), color=DARK_BLUE)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, color=color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.167


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "jiuan-lite 训推平台 | 项目成果与演示汇报"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in hp.runs:
            set_run_font(run, size=8.5, bold=True, color=MUTED)
        hp.paragraph_format.space_after = Pt(2)
        p_pr = hp._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), "D9E2F3")
        p_bdr.append(bottom)
        p_pr.append(p_bdr)
        add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("jiuan-lite 训推平台")
    set_run_font(run, size=26, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run("全项目跑通演示与阶段成果汇报")
    set_run_font(run, size=17, bold=True, color=DARK_BLUE)

    table = doc.add_table(rows=4, cols=2)
    data = [
        ("日期", "2026-07-20"),
        ("演示对象", "老板 / 项目负责人"),
        ("项目定位", "通用语言模型训练闭环 demo / SFT 验证平台"),
        ("演示入口", "http://127.0.0.1:8000/?v=custom9"),
    ]
    for row, (label, value) in zip(table.rows, data):
        row.cells[0].text = label
        row.cells[1].text = value
        shade_cell(row.cells[0], LIGHT)
        for idx, cell in enumerate(row.cells):
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    set_run_font(run, size=10, bold=(idx == 0), color=INK)
    set_table_geometry(table, [1900, 7460])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("汇报摘要")
    set_run_font(run, size=12, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(0)
    add_inline(
        p,
        "已完成对标久安核心方法论的轻量训推平台 PoC：覆盖数据、训练、推理、评测、标注回灌、RAG 与 Agent 自动迭代；最新版本增加模型身份门禁、Judge 可靠性状态和可选择历史模型的 LoRA 续训。",
        base_size=11,
    )
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EEF4FB")
    p_pr.append(shd)
    doc.add_page_break()


def parse_table(lines, start):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [c.strip() for c in lines[idx].strip().strip("|").split("|")]
        rows.append(cells)
        idx += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows, idx


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    font_size = 8.4 if cols <= 4 else 7.4
    for r_idx, source_row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            text = source_row[c_idx] if c_idx < len(source_row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, text, base_size=font_size, color=INK)
            if r_idx == 0:
                shade_cell(cell, LIGHT)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    set_repeat_header(table.rows[0])
    set_table_geometry(table, widths_for(cols))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def build(md_path: Path, out_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_header_footer(doc)
    add_cover(doc)

    idx = 0
    in_code = False
    code_lang = ""
    code_lines = []
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline(p, " ".join(x.strip() for x in paragraph_buffer), base_size=10.5)
            paragraph_buffer = []

    while idx < len(lines):
        raw = lines[idx]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if not in_code:
                in_code = True
                code_lang = stripped[3:].strip()
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.15)
                p.paragraph_format.right_indent = Inches(0.15)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                text = "\n".join(code_lines)
                if code_lang:
                    text = f"[{code_lang}]\n{text}"
                run = p.add_run(text)
                set_run_font(run, ascii_name="Consolas", east_asia="Microsoft YaHei", size=8.3, color=INK)
                p_pr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), CODE_FILL)
                p_pr.append(shd)
                in_code = False
                code_lines = []
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if not stripped:
            flush_paragraph()
            idx += 1
            continue
        if stripped == "---":
            flush_paragraph()
            idx += 1
            continue
        if stripped.startswith("|") and idx + 1 < len(lines) and lines[idx + 1].strip().startswith("|"):
            flush_paragraph()
            rows, idx = parse_table(lines, idx)
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            if level == 1:
                idx += 1
                continue
            p = doc.add_paragraph(style=f"Heading {level - 1}")
            add_inline(p, heading.group(2), base_size=16 if level == 2 else 13, color=BLUE)
            idx += 1
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.right_indent = Inches(0.12)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, stripped.lstrip("> "), base_size=10.3, color=DARK_BLUE)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF4FB")
            p_pr.append(shd)
            idx += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        number = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or number:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(p, (bullet or number).group(1), base_size=10.5)
            idx += 1
            continue
        paragraph_buffer.append(stripped)
        idx += 1

    flush_paragraph()
    core = doc.core_properties
    core.title = "jiuan-lite 训推平台全项目跑通演示与阶段成果汇报"
    core.subject = "通用语言模型训练闭环 demo / SFT 验证平台"
    core.author = "jiuan-lite 项目组"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_boss_demo_docx.py input.md output.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
